"""实际生成 DOCX/PDF 样例；OCR 仅模拟识别接口，不要求 Tesseract。"""
from io import BytesIO
from types import SimpleNamespace as N
from unittest import TestCase
from unittest.mock import Mock, patch
from docx import Document
from pypdf import PdfWriter
from PIL import Image
import document_parser as parser
import ocr_utils


class TestDocumentParsing(TestCase):
    def test_utf8_bom_and_markdown(self):
        for suffix in (".txt", ".md"):
            self.assertEqual(parser.parse_document(b"\xef\xbb\xbf" + "中文".encode(), suffix), "中文")

    def test_invalid_encoding_and_unknown_format(self):
        with self.assertRaises(UnicodeDecodeError):
            parser.parse_document(b"\xff\xfe", ".txt")
        with self.assertRaises(ValueError):
            parser.parse_document(b"data", ".exe")

    def test_docx_preserves_paragraph_table_order(self):
        document = Document()
        document.add_paragraph("表前")
        table = document.add_table(rows=1, cols=2)
        table.cell(0, 0).text = "左格"
        table.cell(0, 1).text = "右格"
        document.add_paragraph("表后")
        stream = BytesIO()
        document.save(stream)
        text = parser.parse_document(stream.getvalue(), ".docx")
        self.assertLess(text.index("表前"), text.index("左格\t右格"))
        self.assertLess(text.index("右格"), text.index("表后"))

    def test_text_pdf_pages_not_discarded(self):
        reader = N(is_encrypted=False, pages=[N(extract_text=Mock(return_value="第一页")), N(extract_text=Mock(return_value="第二页"))])
        with patch.object(parser, "PdfReader", return_value=reader), patch.object(parser, "extract_pdf_page_text") as ocr:
            text = parser.parse_document(b"mock", ".pdf")
        self.assertIn("## 第 1 页", text)
        self.assertIn("第一页", text)
        self.assertLess(text.index("第一页"), text.index("第二页"))
        ocr.assert_not_called()

    def test_mixed_pdf_only_ocr_blank_text_page(self):
        reader = N(is_encrypted=False, pages=[N(extract_text=Mock(return_value="文本页")), N(extract_text=Mock(return_value=" "))])
        with patch.object(parser, "PdfReader", return_value=reader), patch.object(parser, "extract_pdf_page_text", return_value="扫描文字") as ocr:
            text = parser.parse_document(b"mock", ".pdf")
        ocr.assert_called_once_with(b"mock", 1)
        self.assertIn("OCR", text)
        self.assertLess(text.index("文本页"), text.index("扫描文字"))

    def test_unreadable_pdf_page_reports_number(self):
        reader = N(is_encrypted=False, pages=[N(extract_text=Mock(return_value=None))])
        with patch.object(parser, "PdfReader", return_value=reader), patch.object(parser, "extract_pdf_page_text", return_value=""):
            with self.assertRaisesRegex(ValueError, "第 1 页"):
                parser.parse_document(b"mock", ".pdf")

    def test_pdf_page_limit_and_encryption(self):
        reader = N(is_encrypted=False, pages=[None] * 101)
        with patch.object(parser, "PdfReader", return_value=reader):
            with self.assertRaisesRegex(ValueError, "100"):
                parser.parse_document(b"mock", ".pdf")
        writer = PdfWriter()
        writer.add_blank_page(width=72, height=72)
        writer.encrypt("test-password")
        stream = BytesIO()
        writer.write(stream)
        with self.assertRaisesRegex(ValueError, "加密"):
            parser.parse_document(stream.getvalue(), ".pdf")

    def test_real_blank_pdf_uses_ocr(self):
        writer, stream = PdfWriter(), BytesIO()
        writer.add_blank_page(width=72, height=72)
        writer.write(stream)
        with patch.object(parser, "extract_pdf_page_text", return_value="模拟扫描文字") as ocr:
            self.assertIn("模拟扫描文字", parser.parse_document(stream.getvalue(), ".pdf"))
        self.assertEqual(ocr.call_args.args[1], 0)

    def test_image_types_dispatch_to_ocr(self):
        for suffix in (".png", ".jpg", ".jpeg"):
            with self.subTest(suffix=suffix), patch.object(parser, "extract_image_text", return_value="图片文字") as ocr:
                self.assertEqual(parser.parse_document(b"mock", suffix), "图片文字")
                ocr.assert_called_once_with(b"mock")

    def test_damaged_binary_files_raise(self):
        for suffix in (".pdf", ".docx"):
            with self.subTest(suffix=suffix), self.assertRaises(Exception):
                parser.parse_document(b"not a document", suffix)

    def test_ocr_language_timeout_and_image_cleanup(self):
        with Image.new("RGB", (20, 20), "white") as image:
            with patch.object(ocr_utils.pytesseract, "image_to_string", return_value="  识别结果  ") as recognize:
                self.assertEqual(ocr_utils.recognize_image(image), "识别结果")
        self.assertEqual(recognize.call_args.kwargs, {"lang": "chi_sim+eng", "timeout": 30})

